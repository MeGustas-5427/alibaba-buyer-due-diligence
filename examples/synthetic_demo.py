"""Generate an entirely invented batch. Never use these observations as research.

The interactive demo stops at the PDF visual-review gate. It does not synthesize
a visual approval. Unit tests may separately simulate that gate in temp folders.
"""
import argparse
import copy
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]/"scripts"))
from extract_research_view import HEADERS, atomic, cell_text, digest, now, project, read_json, require
from validate_report import DIMENSIONS, RISKS, SCORES, draft
import workflow


def source(count=3):
    records = []
    for i in range(count):
        crm, login = f"synthetic-crm-{i:02d}", f"synthetic-login-{i:02d}"
        company = f"SYNTHETIC {i+1:02d} 虚构海湾设备"
        if i == 0:
            company += " International Distribution and Service Company"
        records.append({"crmCustomerId": crm, "buyerLoginId": login,
            "buyerList": {"customerId": crm, "customer": {"customerId": crm, "loginId": login}},
            "customerDetail": {"customerId": crm, "companyName": company, "fieldsMissingFromList": {"ownerName": "虚构业务员甲", "contacts": [{"contactName": f"虚构联系人{i+1}", "position": "采购"}], "address": {"country": "GB"}}},
            "buyerHomepageStatus": "collected", "buyerHomepage": {"customerId": crm,
                "buyerProfile": {"memberId": login, "registerYear": 2020, "registeredCountryRegion": "GB"},
                "enterpriseInfo": {"companyName": company, "officialWebsite": f"https://example.com/synthetic/{i}", "miniSiteUrl": "", "businessTypes": [{"id": 2003}], "address": {"region": "GB", "addressText": "SYNTHETIC Harbour Test Zone"}},
                "purchasePreference": {"industryPreference": [{"label": "虚构设备测试品类"}]},
                "activity90d": {"loginDays": 0, "productViewCount": 8, "validInquiryCount": 2}, "onlineTrade": {"totalOrderCount": 0, "totalOrderVolumeUsd": 0}}})
    return {"schemaVersion": "nanyue.due-diligence-input.v3", "generatedAt": "2026-10-05T00:00:00+00:00", "sourceDate": "2026-10-04", "readerInstructions": ["SYNTHETIC TEST ONLY"], "source": {"synthetic": True}, "records": records}


def workbook(path, payload, batch_size=10):
    from openpyxl import Workbook
    view = project(payload, batch_size)
    book = Workbook()
    book.active.title = "背调输入"
    batches = (len(view["records"])+batch_size-1)//batch_size
    sheets = {"背调输入": view["records"], **{f"批次{b:02d}": [r for r in view["records"] if r["batchNo"] == b] for b in range(1, batches+1)}}
    for name, rows in sheets.items():
        sheet = book[name] if name in book.sheetnames else book.create_sheet(name)
        sheet.append([label for _, label in HEADERS])
        for row in rows:
            sheet.append([cell_text(row[key]) for key, _ in HEADERS])
    meta = book.create_sheet("生成说明")
    for row in [("项目", "值"), ("schemaVersion", "nanyue.due-diligence-research-view.v1"), ("generatedAt", payload["generatedAt"]), ("sourceDate", payload["sourceDate"]), ("batchSize", batch_size), ("batchCount", batches)]:
        meta.append(row)
    book.save(path)
    book.close()


def prepare(base, count=3, session=None):
    base = Path(base)
    base.mkdir(parents=True, exist_ok=True)
    payload = source(count)
    atomic(base/"synthetic-input.json", payload)
    workbook(base/"synthetic-input.xlsx", payload)
    run = Path(workflow.initialize(base/"synthetic-input.json", base/"synthetic-input.xlsx", base, session=session)["run"])
    workflow.advance(run)
    workflow.advance(run)
    return run


def fill(run, matched=2):
    """Test fixture, intentionally refuses any non-SYNTHETIC customer."""
    state = workflow.load_state(run)
    view = workflow.inputs(state)
    require(all(r["companyName"].startswith("SYNTHETIC ") for r in view["records"]), "synthetic fixture refuses real customer data")
    data = draft(view, state["inputs"]["json"]["sha256"])
    for i, record in enumerate(data["records"]):
        filename = f"evidence/synthetic-{i:02d}.md"
        atomic(Path(run)/filename, f"# SYNTHETIC observation {i}\n\nOffline, wholly invented fixture. NO web search occurred.\nCompany and country match only within this test universe.\nSanctions download deliberately unavailable in this offline test.\n")
        record["attempts"] = [{"kind": k, "language": "en", "query": record["companyName"]+" SYNTHETIC "+k,
            "checkedAt": now(), "status": "limited" if k == "compliance" else "succeeded", "detail": "SYNTHETIC offline observation; not a real search result",
            "evidenceFile": filename, "evidenceSha256": digest(Path(run)/filename)} for k in ("identity", *RISKS)]
        record["sources"] = [{"url": f"https://{domain}/synthetic/{i}", "title": "SYNTHETIC company fixture", "publisher": f"SYNTHETIC Publisher {j}", "publisherGroup": f"synthetic-publisher-{j}",
            "accessedAt": now(), "supports": ["identity", *DIMENSIONS, *RISKS], "attemptIndex": 0} for j, domain in enumerate(("example.com", "example.org"))]
        record["riskChecks"] = {k: {"status": "limited" if k == "compliance" else "passed", "summary": "SYNTHETIC offline fixture; no real-world conclusion", "attemptIndexes": [j+1], "sourceIndexes": [] if k == "compliance" else [0]} for j, k in enumerate(RISKS)}
        record["screening"] = {"status": "limited", "reason": "SYNTHETIC offline scenario: full UK list not downloaded; never zero hits", "termsFile": None, "receipt": None}
        record["matchStatus"] = "matched" if i < matched else "unconfirmed"
        record["identificationStrength"] = "high" if i < matched else "low"
        record["matchReason"] = "SYNTHETIC only: company and country anchors agree" if i < matched else "SYNTHETIC ambiguous company; do not attribute candidate facts"
        record["identityEvidence"] = [{"field": f, "sourceIndexes": [0, 1], "reason": "SYNTHETIC agreement"} for f in ("companyName", "countryRegion")] if i < matched else []
        record["publicSummary"] = "全虚构演示：设备分销企业，英文长名称用于版式回归。此描述不是现实公司事实。"
        record["riskJudgement"] = "全虚构演示：名单未完成读取，合规待核验；高风险客户也保留。"
        record["followup"] = "先核验主体；补充采购规格与预算；确认付款及交付条件。"
        record["dimensions"] = {k: {"summary": "全虚构演示；缺少实际公开证据，不对现实主体下结论。", "evidenceLevel": "insufficient", "sourceIndexes": []} for k in DIMENSIONS}
        record["scores"] = {k: {"value": 8.8 if k == "risk" else None, "reason": "全虚构风险高分测试；其他分项无真实依据，不补零。", "sourceIndexes": [0] if k == "risk" else []} for k in SCORES} if i < matched else {}
    atomic(Path(run)/"research.json", data)
    return data


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", required=True, type=Path)
    parser.add_argument("--customers", type=int, default=9)
    parser.add_argument("--matched", type=int, default=8)
    args = parser.parse_args()
    require(not args.out.exists(), "demo output must be a new directory")
    require(0 <= args.matched <= args.customers <= 30, "invalid demo counts")
    run = prepare(args.out, args.customers)
    fill(run, args.matched)
    for _ in range(3):
        workflow.advance(run)
    try:
        workflow.advance(run)
    except Invalid as exc:
        if "Required file missing: visual-review.json" not in str(exc):
            raise
    if not args.matched:
        workflow.advance(run, finalize=True)
    print(encode(workflow.status(run)))


if __name__ == "__main__":
    from extract_research_view import Invalid, encode
    main()
