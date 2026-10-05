"""Standalone, read-only Alibaba v3 / 26-column workbook contract.

No collector import, network access, or data upload. Field mappings implement
the existing export format; see references/input-contract.md.
"""
from __future__ import annotations

import hashlib
import ipaddress
import json
import math
import os
import re
import tempfile
from datetime import date, datetime, timezone
from pathlib import Path
from urllib.parse import parse_qsl, unquote, urlsplit, urlunsplit

INPUT_SCHEMA = "nanyue.due-diligence-input.v3"
VIEW_SCHEMA = "alibaba.buyer-due-diligence.research-view.v1"
HEADERS = list(zip(
    "batchNo rowNo crmCustomerId buyerLoginId companyName contactName salesperson countryRegion registerYear officialWebsite miniSiteUrl enterpriseAddress businessType position industryPreference activityLoginDays90d activityProductViews90d activitySearches90d validInquiryCount90d repliedInquiryCount90d validRfqCount90d quotationReceivedCount90d quotationReadCount90d onlineTradeOrderCount onlineTradeVolumeUsd searchHints".split(),
    "批次 序号 CRM客户ID 买家登录ID 公司名称 联系人 负责业务员 国家/地区 Alibaba账号注册年份 官网 阿里MiniSite 企业地址 经营类型 职位 行业偏好 90天登录天数 90天商品浏览量 90天搜索次数 90天有效询盘 90天已回复询盘 90天RFQ 90天收到报价 90天已读报价 在线交易数量 在线交易金额USD 检索线索".split()))
DISPLAY = ("companyName", "contactName", "salesperson", "countryRegion")
TYPE_IDS = dict(zip(map(str, range(2000, 2009)),
                   "线上零售商 工厂 贸易公司 批发/分销商 线下零售商 买办/采购代理 线上零售商 散客 其他".split()))
TYPE_NAMES = dict(zip("manufacturerOrFactory tradingCompany distributorWholesaler retailer buyerOffice onlineShop soho other".split(),
                     "工厂 贸易公司 批发/分销商 线下零售商 买办/采购代理 线上零售商 散客 其他".split()))
TRANSPORT = set("authorization cookie cookies credentials headers html password proxyauthorization rawhtml requestheaders secret session sessionid setcookie tbtoken token".split())
PRIVATE = TRANSPORT | set("email emailaddress emailfull emailverified fullemail maskedemail phone phonenumber mobile mobilenumber contactdetails buyerprofilelink profileurl crmcustomerid buyerloginid".split())
EMAIL = re.compile(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}")
SECRET = re.compile(r"(?i)(?:token|authorization|cookie|password|secret|api[_ -]?key|session[_ -]?id)\s*[:=]\s*\S+|\bbearer\s+[\w.~-]+|\beyJ[\w-]+\.[\w-]+\.[\w-]+")
HTML = re.compile(r"(?i)<(?:!doctype|/?(?:html|head|body|script|style|div|span|table))\b")


class Invalid(ValueError):
    """A non-sensitive, actionable validation failure."""


def require(ok, message):
    if not ok:
        raise Invalid(message)


def exact(value, keys, label):
    require(isinstance(value, dict) and set(value) == set(keys), label + ": wrong fields")


def now():
    return datetime.now(timezone.utc).isoformat()


def timestamp(value):
    require(isinstance(value, str), "timestamp must be text")
    try:
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
        require(parsed.tzinfo is not None, "timestamp needs timezone")
    except ValueError as exc:
        raise Invalid("invalid timestamp") from exc
    return parsed


def digest(path):
    with Path(path).open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def encode(value):
    return json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + "\n"


def read_json(path):
    def unique(pairs):
        result = {}
        for key, value in pairs:
            require(key not in result, "JSON contains duplicate keys")
            result[key] = value
        return result
    return json.loads(Path(path).read_text(encoding="utf-8-sig"), object_pairs_hook=unique,
                      parse_constant=lambda _: (_ for _ in ()).throw(Invalid("nonfinite JSON number")))


def atomic(path, value):
    path = Path(path)
    require(not path.is_symlink(), "refusing symlink output")
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temp = tempfile.mkstemp(prefix=".write-", dir=path.parent)
    try:
        with os.fdopen(fd, "w", encoding="utf-8", newline="\n") as stream:
            stream.write(value if isinstance(value, str) else encode(value))
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temp, path)
    finally:
        Path(temp).unlink(missing_ok=True)


def public_url(value):
    require(isinstance(value, str) and value == value.strip() and not re.search(r"[\s\\]", value), "invalid public URL")
    try:
        parts = urlsplit(value)
        host = (parts.hostname or "").lower().rstrip(".")
        port = parts.port
    except ValueError as exc:
        raise Invalid("invalid public URL") from exc
    require(parts.scheme in {"https", "http"} and "." in host and not parts.username and not parts.password, "URL must be public HTTP(S)")
    require(not host.endswith((".local", ".internal", ".localhost")) and host not in {"localhost", "profile.alibaba.com", "i.alibaba.com"}, "private URL prohibited")
    try:
        address = ipaddress.ip_address(host)
    except ValueError:
        address = None
    require(address is None or address.is_global, "private IP URL prohibited")
    require(address is not None or not re.fullmatch(r"[\d.]+", host), "ambiguous numeric URL host")
    for key, _ in parse_qsl(parts.query + "&" + parts.fragment, keep_blank_values=True):
        key = re.sub(r"[^a-z0-9]", "", key.lower())
        require(not any(t in key for t in ("token", "credential", "signature", "session", "secret", "password", "auth", "apikey")) and key not in {"sig", "policy", "keypairid", "phone", "mobile", "email", "tel"}, "credential/contact URL prohibited")
    require(not SECRET.search(unquote(value)) and not EMAIL.search(unquote(value)), "sensitive URL prohibited")
    require(not re.search(r"(?:\+|%2b)\d[\d .()-]{6,}\d", unquote(parts.query).lower()), "phone-like URL query prohibited")
    require(not any(t in unquote(parts.path).lower() for t in ("my_profile", "myprofile", "buyer_profile")), "private profile prohibited")
    netloc = host + (f":{port}" if port and (parts.scheme, port) not in {("https", 443), ("http", 80)} else "")
    # Tracking and fragments do not make independent evidence.
    return urlunsplit((parts.scheme, netloc, parts.path.rstrip("/") or "/", "", ""))


def safe_text(value, label="text", phone=True):
    require(isinstance(value, str), label + ": expected text")
    require(not EMAIL.search(value) and not SECRET.search(value) and not HTML.search(value), label + ": sensitive content prohibited")
    for found in re.findall(r"https?://[^\s<>\]\)\"']+", value):
        public_url(found.rstrip(".,;，。；"))
    if phone:
        without_dates = re.sub(r"\b\d{4}-\d{2}-\d{2}(?:T[\d:.+Z-]+)?\b", "", value)
        for candidate in re.findall(r"(?<!\w)\+?\d[\d ().-]{5,}\d(?!\w)", without_dates):
            digits = re.sub(r"\D", "", candidate)
            require(not (7 <= len(digits) <= 15 and (len(digits) >= 10 or any(c in candidate for c in "+() "))), label + ": phone-like content prohibited")
    return value


def safety_walk(value, private=False, key=""):
    if isinstance(value, dict):
        for name, item in value.items():
            normalized = re.sub(r"[^a-z0-9]", "", name.lower())
            require(normalized not in (TRANSPORT if private else PRIVATE), "prohibited field")
            safety_walk(item, private, name)
    elif isinstance(value, list):
        for item in value:
            safety_walk(item, private, key)
    elif isinstance(value, str):
        if private and key.lower() in {"buyerprofilelink", "profileurl"}:
            require(not HTML.search(value), "raw HTML prohibited")
        elif private:
            require(not SECRET.search(value) and not HTML.search(value), "source contains credentials/HTML")
        else:
            safe_text(value, key, phone=key not in {"inputSha256", "sha256", "termsSha256", "pdfSha256", "sourceSha256", "accessedAt", "checkedAt", "createdAt", "url", "evidenceSha256"})


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, (int, float)):
        require(math.isfinite(value), "nonfinite input")
        return value
    if isinstance(value, list):
        return ", ".join(str(cell_text(v)) for v in value if cell_text(v) != "")
    if isinstance(value, dict):
        for key in ("label", "name", "displayName", "displayText", "other", "id", "code", "mcmsKey"):
            if value.get(key) not in (None, ""):
                return cell_text(value[key])
        return "" if all(v in (None, "", [], {}) for v in value.values()) else json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return str(value).strip()


def business_type(enterprise):
    def translated(value):
        value = str(value)
        return TYPE_IDS.get(value, TYPE_NAMES.get(value.removeprefix("businessType@"), value))
    raw = enterprise.get("businessTypeText")
    if raw not in (None, ""):
        values = [translated(v.strip()) for v in str(raw).replace("，", ",").split(",") if v.strip()]
    else:
        values = []
        for option in enterprise.get("businessTypes") or []:
            if isinstance(option, dict):
                values.append(cell_text(option.get("label")) or TYPE_NAMES.get(str(option.get("mcmsKey", "")).removeprefix("businessType@"), TYPE_IDS.get(str(option.get("id")), cell_text(option.get("mcmsKey") or option.get("id")))))
            else:
                values.append(cell_text(option))
    return ", ".join(dict.fromkeys(str(v) for v in values if v not in (None, "")))


def project(payload, batch_size=10):
    require(isinstance(payload, dict) and payload.get("schemaVersion") == INPUT_SCHEMA, "expected Alibaba daily v3 input")
    require(type(batch_size) is int and batch_size > 0, "invalid batch size")
    source_date = payload.get("sourceDate")
    require(isinstance(source_date, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", source_date), "sourceDate must be YYYY-MM-DD")
    date.fromisoformat(source_date)
    timestamp(payload.get("generatedAt"))
    records = payload.get("records")
    require(isinstance(records, list), "records must be an array")
    safety_walk(payload, private=True)
    rows, seen_crm, seen_login = [], set(), set()
    for index, item in enumerate(records, 1):
        exact(item, ("crmCustomerId", "buyerLoginId", "buyerList", "customerDetail", "buyerHomepageStatus", "buyerHomepage"), f"input record {index}")
        crm, login = item["crmCustomerId"], item["buyerLoginId"]
        for value, seen in ((crm, seen_crm), (login, seen_login)):
            require(isinstance(value, str) and value.strip() == value and value and value not in seen, f"record {index}: missing/duplicate/noncanonical identity")
            seen.add(value)
        require(not login.endswith((",", "，")), "noncanonical login")
        listing, detail, profile = item["buyerList"], item["customerDetail"], item["buyerHomepage"]
        require(isinstance(listing, dict) and isinstance(detail, dict), "list/detail must be objects")
        status = item["buyerHomepageStatus"]
        require((status == "collected" and isinstance(profile, dict)) or (status == "not_collected_no_profile_link" and profile is None), "homepage status mismatch")
        profile = profile or {}
        customer, buyer = listing.get("customer") or {}, profile.get("buyerProfile") or {}
        for obj in (listing, customer, detail, profile):
            require(obj.get("customerId") in (None, "") or str(obj["customerId"]).strip() == crm, "nested CRM identity mismatch")
        for value in (customer.get("loginId"), buyer.get("memberId")):
            require(value in (None, "") or str(value).strip().removesuffix(",").strip() == login, "nested login identity mismatch")
        fields = detail.get("fieldsMissingFromList") or {}
        contact = (fields.get("contacts") or [{}])[0]
        enterprise = profile.get("enterpriseInfo") or {}
        activity, trade = profile.get("activity90d") or {}, profile.get("onlineTrade") or {}
        company = detail.get("companyName") or profile.get("companyName") or enterprise.get("companyName") or ""
        country = buyer.get("registeredCountryRegion") or (buyer.get("contactAddress") or {}).get("countryRegion") or (enterprise.get("address") or {}).get("region") or (fields.get("address") or {}).get("country") or ""
        website, mini, kind = enterprise.get("officialWebsite") or fields.get("website") or "", enterprise.get("miniSiteUrl") or "", business_type(enterprise)
        values = [((index-1)//batch_size)+1, index, crm, login, company,
                  contact.get("contactName") or buyer.get("displayName") or "", fields.get("ownerName") or "", country,
                  buyer.get("registerYear"), website, mini, (enterprise.get("address") or {}).get("addressText") or fields.get("addressText") or "", kind,
                  enterprise.get("position") or contact.get("position"), cell_text((profile.get("purchasePreference") or {}).get("industryPreference")),
                  *[activity.get(k) for k in ("loginDays", "productViewCount", "searchCount", "validInquiryCount", "repliedInquiryCount", "validRfqCount", "quotationReceivedCount", "quotationReadCount")],
                  trade.get("totalOrderCount"), trade.get("totalOrderVolumeUsd"), "; ".join(str(v) for v in (company, country, website, mini, kind) if v)]
        row = dict(zip((k for k, _ in HEADERS), values, strict=True))
        for key, value in row.items():
            if isinstance(value, str) and key not in {"crmCustomerId", "buyerLoginId"}:
                safe_text(value, f"record {index}.{key}")
        require(all(isinstance(row[k], str) for k in DISPLAY), "display fields must be strings")
        rows.append(row)
    return {"schemaVersion": VIEW_SCHEMA, "sourceDate": source_date, "records": rows}


def validate_pair(input_path, xlsx_path, phase_path=None):
    from openpyxl import load_workbook
    payload = read_json(input_path)
    workbook = load_workbook(xlsx_path, read_only=True, data_only=False)
    try:
        require("生成说明" in workbook.sheetnames and "背调输入" in workbook.sheetnames, "workbook sheets missing")
        meta = dict(list(workbook["生成说明"].values)[1:])
        require(meta.get("sourceDate") == payload.get("sourceDate"), "workbook date mismatch")
        require(meta.get("generatedAt") == payload.get("generatedAt"), "workbook generation timestamp mismatch")
        require(meta.get("schemaVersion") == "nanyue.due-diligence-research-view.v1", "workbook schema mismatch")
        view = project(payload, meta.get("batchSize"))
        rows = view["records"]
        batches = (len(rows) + meta["batchSize"] - 1) // meta["batchSize"]
        require(meta.get("batchCount") == batches, "workbook batch count mismatch")
        expected = {"背调输入": rows, **{f"批次{i:02d}": [r for r in rows if r["batchNo"] == i] for i in range(1, batches+1)}}
        require(set(workbook.sheetnames) == set(expected) | {"生成说明"}, "unexpected workbook sheets")
        for name, data in expected.items():
            actual = list(workbook[name].values)
            wanted = [[label for _, label in HEADERS], *[[cell_text(row[k]) for k, _ in HEADERS] for row in data]]
            require(len(actual) == len(wanted), name + ": row count mismatch")
            for i, (left, right) in enumerate(zip(actual, wanted)):
                require(["" if v is None else v for v in left] == right, f"{name} row {i}: 26-field mismatch")
            require(not any(c.data_type == "f" for row in workbook[name] for c in row), "workbook formulas prohibited")
    finally:
        workbook.close()
    if phase_path:
        phase = read_json(phase_path)
        final = phase.get("phases", {}).get("final_validated", {})
        summary = final.get("summary", {})
        require(phase.get("currentPhase") == phase.get("lastCheckedPhase") == "final_validated" and phase.get("lastStatus") == final.get("status") == "complete" and final.get("errors") == [], "upstream phase not final_validated")
        require(summary.get("sourceDate") == payload["sourceDate"] and summary.get("nanyueRecords") == len(rows), "upstream phase scope mismatch")
        for key, path in (("nanyueJson", input_path), ("nanyueWorkbook", xlsx_path)):
            require(Path(summary.get(key, "")).resolve() == Path(path).resolve(), "upstream phase path mismatch")
    return view
